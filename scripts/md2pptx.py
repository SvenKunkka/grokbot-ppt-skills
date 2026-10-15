#!/usr/bin/env python3
"""Markdown 大纲 -> 可编辑 PPTX（16:9）。# 新页标题，- 要点，> 备注。"""
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt

def parse(text):
    slides, cur = [], None
    for line in text.splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            cur = {"title": s[2:].strip(), "bullets": [], "notes": []}; slides.append(cur)
        elif cur is None or not s.strip():
            continue
        elif s.lstrip().startswith(("- ", "* ")):
            lvl = (len(s) - len(s.lstrip())) // 2
            cur["bullets"].append((lvl, s.lstrip()[2:].strip()))
        elif s.startswith("> "):
            cur["notes"].append(s[2:].strip())
        else:
            cur["bullets"].append((0, s.strip()))
    return slides

def build(slides, out):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, sd in enumerate(slides):
        layout = prs.slide_layouts[0 if i == 0 and not sd["bullets"] else 1]
        sl = prs.slides.add_slide(layout); sl.shapes.title.text = sd["title"]
        if sd["bullets"] and len(sl.placeholders) > 1:
            tf = sl.placeholders[1].text_frame; tf.clear()
            for j, (lvl, t) in enumerate(sd["bullets"]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text, p.level = t, min(lvl, 4); p.font.size = Pt(20 - 2 * min(lvl, 3))
        if sd["notes"]:
            sl.notes_slide.notes_text_frame.text = "\n".join(sd["notes"])
    prs.save(out)

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("md"); ap.add_argument("-o", "--out", default="out.pptx")
    a = ap.parse_args(); s = parse(open(a.md, encoding="utf-8").read())
    if not s: raise SystemExit("大纲里没有以 # 开头的页面标题")
    build(s, a.out); print(f"已生成 {a.out}，共 {len(s)} 页")
